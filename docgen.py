# -*- coding: utf-8 -*-
"""Генерация претензий и исковых заявлений (.docx).

Вёрстка и формулировки повторяют реальные образцы компании:
- иск в Экономический суд      — «образец иска в экономический суд»
- иск в Арбитражную комиссию   — «образец иска в арбитражную комиссию», «образец иска (одна накладная)»,
                                 «образец иска (со списанием налоговой)»,
                                 «образец иска с частичной оплатой» (частичная оплата)
- претензия (для обоих типов)  — «исправ_претензия.docx»

Что взято из образцов: поля страницы, шапка бланка (16 пт прописными + 9 пт
реквизиты), адресный блок в правой половине листа, основной текст 14 пт по
ширине с красной строкой 1,25 см, суммы полужирным, цитаты из договора и
кодекса курсивом, состав и порядок приложений, подпись.
"""
import datetime as dt
import os
import re
from dataclasses import dataclass

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn

from company import (
    SELLER,
    DEFAULT_PENALTY_RATE,
    DISPUTE_BODY,
    CONTRACT_CLAUSES,
    ATTACHMENTS,
    PAYMENT_TERM,
    CHARGE_INTEREST_366,
    CHARGE_INTEREST_366_IN_PRETENZIA,
)
from parser import DebtLine
from rates import load_refi_rates, interest_366, state_duty
import payments as pay

FONT = "Times New Roman"
BODY_PT = 14        # основной текст
HEAD_PT = 16        # название компании в шапке
SMALL_PT = 9        # реквизиты в шапке
BODY_INDENT = 1.25  # красная строка, см
ADDR_INDENT = 7.0   # адресный блок — правая половина листа, см

STATUTE_OF_LIMITATIONS_DAYS = 1095  # ~3 года, ст. 197 ГК РБ

RU_MONTHS = [
    "", "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]

# Типы документов в отчёте 1С названы по-своему — приводим к языку суда.
DOC_TYPE_NAMES = {
    "Руб.накл": ("товарно-транспортная накладная", "ТН", "ТТН"),
    "Накл.материалы": ("товарная накладная", "ТН", "ТТН"),
    "Счет за услуги": ("счёт за услуги", "счёт", "счёт"),
}


def doc_full(t: str) -> str:
    return DOC_TYPE_NAMES.get(t, (t, t, t))[0]


def doc_short(t: str) -> str:
    return DOC_TYPE_NAMES.get(t, (t, t, t))[1]


def doc_short_ex(t: str) -> str:
    """Короткое имя в биржевых документах — там пишут «ТТН»."""
    return DOC_TYPE_NAMES.get(t, (t, t, t))[2]


# ---------------------------------------------------------------------------
# Форматирование
# ---------------------------------------------------------------------------

def ru_date(d: dt.date) -> str:
    return f"{d.day:02d} {RU_MONTHS[d.month]} {d.year}г."


def contract_date_words(short: str) -> str:
    """'28.05.26' -> '28 мая 2026г.'; если разобрать не вышло — вернёт как есть."""
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$", (short or "").strip())
    if not m:
        return short
    day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if year < 100:
        year += 2000
    if not 1 <= month <= 12:
        return short
    return f"{day} {RU_MONTHS[month]} {year}г."


def contract_date_full(short: str) -> str:
    """'28.05.26' -> '28.05.2026'."""
    m = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$", (short or "").strip())
    if not m:
        return short
    year = int(m.group(3))
    if year < 100:
        year += 2000
    return f"{int(m.group(1)):02d}.{int(m.group(2)):02d}.{year}"


def fmt_money(x: float) -> str:
    """17044.02 -> '17 044,02' (неразрывный пробел, как в образцах)."""
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def fmt_rate(x: float) -> str:
    return f"{x:g}".replace(".", ",")


def _setup(doc: Document, left=2.22, right=1.25, top=1.0, bottom=1.25):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = style.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(left), Cm(right)
        s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)


def _run(par, text, bold=False, italic=False, underline=False, size=BODY_PT):
    r = par.add_run(text)
    r.bold, r.italic, r.underline = bold, italic, underline
    r.font.size = Pt(size)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:cs"), FONT)
    return r


def _par(doc, align=None, first_line=None, left=None, space_after=0, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    if left is not None:
        pf.left_indent = Cm(left)
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    return p


def _line(doc, text, align=None, first_line=None, left=None, bold=False,
          italic=False, size=BODY_PT, space_after=0):
    """Абзац из одного куска текста."""
    p = _par(doc, align=align, first_line=first_line, left=left, space_after=space_after)
    _run(p, text, bold=bold, italic=italic, size=size)
    return p


def _body(doc, parts, space_after=0):
    """Абзац основного текста: по ширине, с красной строки.

    parts — строка либо список кусков: str или (текст, {'bold':True} / {'italic':True}).
    """
    p = _par(doc, align=AL.JUSTIFY, first_line=BODY_INDENT, space_after=space_after)
    if isinstance(parts, str):
        parts = [parts]
    for part in parts:
        if isinstance(part, str):
            _run(p, part)
        else:
            text, opts = part
            _run(p, text, **opts)
    return p


def _empty(doc):
    _par(doc)


@dataclass
class BuyerInfo:
    name: str
    name_full: str = ""   # полное наименование с организационно-правовой формой
    address: str = ""
    unp: str = ""
    bank_account: str = ""
    bank_name: str = ""
    bank_code: str = ""
    phone: str = ""
    email: str = ""
    director_fio: str = ""
    director_basis: str = ""


# ---------------------------------------------------------------------------
# Общие блоки
# ---------------------------------------------------------------------------

def _underline(par):
    """Тонкая линия под абзацем — ею в образцах отчёркнута шапка бланка."""
    from docx.oxml import OxmlElement
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    bdr.append(bottom)
    par._p.get_or_add_pPr().append(bdr)


def _letterhead(doc):
    """Шапка бланка: 16 пт прописными по центру + реквизиты 9 пт, снизу линия."""
    _line(doc, SELLER["name_caps"], align=AL.CENTER, bold=True, size=HEAD_PT)
    _line(doc, f'УНП {SELLER["unp_spaced"]}, ОКПО {SELLER["okpo_spaced"]}',
          align=AL.CENTER, size=SMALL_PT)
    _line(doc, f'Юридический адрес: {SELLER["address"]}', size=SMALL_PT)
    bank = _line(doc, f'Реквизиты банка: р/с {SELLER["bank_account"]} в {SELLER["bank_name"]}, '
                      f'код  {SELLER["bank_code"]};', size=SMALL_PT)
    tel = _line(doc, f'тел./ф {SELLER["phone"]}, {SELLER["email"]}', size=SMALL_PT)
    _underline(bank)
    _underline(tel)


def _addr(doc, text, bold=False):
    """Строка адресного блока — правая половина листа, по левому краю."""
    return _line(doc, text, left=ADDR_INDENT, align=AL.LEFT)


def _plaintiff_block(doc):
    _addr(doc, f'Истец: {SELLER["name_full"]},')
    _addr(doc, SELLER["address_short"])
    _addr(doc, f'р/с {SELLER["bank_account"]}')
    _addr(doc, f'в {SELLER["bank_name"]},')
    _addr(doc, f'код банка {SELLER["bank_code"]}.')
    _addr(doc, f'УНП   {SELLER["unp_spaced"]}')
    _addr(doc, f'Тел/ф. {SELLER["phone"]}')


def _defendant_block(doc, buyer: BuyerInfo):
    _addr(doc, f"Ответчик: {buyer.name_full or buyer.name}")
    if buyer.address:
        _addr(doc, buyer.address)
    else:
        _addr(doc, "ЗАПОЛНИТЬ адрес — см. data/contragents.csv")
    if buyer.bank_account:
        _addr(doc, f"р/с {buyer.bank_account}")
    if buyer.bank_name:
        _addr(doc, f"{buyer.bank_name},")
    if buyer.bank_code:
        _addr(doc, f"код банка {buyer.bank_code}")
    if buyer.unp:
        _addr(doc, f"УНП {buyer.unp}")
    if buyer.phone:
        _addr(doc, f"Тел. {buyer.phone}")


def _signature(doc, as_of: dt.date = None, executor=True):
    _empty(doc)
    p = _par(doc)
    _run(p, "Директор")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(11))
    _run(p, "\t" + SELLER["director_fio_initials_first"])
    if as_of:
        _line(doc, as_of.strftime("%d.%m.%Y") + "г.")
    if executor:
        _empty(doc)
        _line(doc, f'Исп. {SELLER["executor_fio"]}', size=SMALL_PT + 2)
        _line(doc, SELLER["executor_phone"], size=SMALL_PT + 2)


ISK_COPY_LINE = "Копия искового заявления для Ответчика"


def _attachments(doc, items):
    """Нумерованный перечень приложений: «N. Текст на N л. в 1 экз.»"""
    _empty(doc)
    _line(doc, "Приложение:", first_line=BODY_INDENT)
    for i, (text, sheets) in enumerate(items, start=1):
        _line(doc, f"{i}. {text} на {sheets} л. в 1 экз.", first_line=BODY_INDENT)


def fix_isk_sheets(paths: list) -> int:
    """Проставляет в «Копия искового заявления для Ответчика на N л.» настоящее
    число страниц — считает его Word. Без Word ничего не меняет. Возвращает,
    сколько документов поправлено."""
    try:
        import pythoncom
        import win32com.client as win32
    except ImportError:
        return 0
    pythoncom.CoInitialize()
    fixed = 0
    try:
        try:
            app = win32.Dispatch("Word.Application")
        except Exception:
            return 0
        app.Visible = False
        app.DisplayAlerts = False
        try:
            for path in paths:
                try:
                    wd = app.Documents.Open(os.path.abspath(path), ReadOnly=True,
                                            AddToRecentFiles=False)
                    pages = int(wd.ComputeStatistics(2))   # wdStatisticPages
                    wd.Close(False)
                except Exception:
                    continue
                try:
                    d = Document(path)
                    for p in d.paragraphs:
                        if ISK_COPY_LINE in p.text and p.runs:
                            p.runs[0].text = re.sub(r"на \S+ л\.", f"на {pages} л.",
                                                    p.runs[0].text)
                            d.save(path)
                            fixed += 1
                            break
                except Exception:
                    continue          # файл занят Word'ом у пользователя — пропускаем
        finally:
            app.Quit()
    finally:
        pythoncom.CoUninitialize()
    return fixed


# ---------------------------------------------------------------------------
# Расчёты
# ---------------------------------------------------------------------------

def _period(line: DebtLine, as_of: dt.date):
    """(первый день просрочки, последний день, дней) с учётом срока давности."""
    first = line.due_date + dt.timedelta(days=1)
    cap = as_of - dt.timedelta(days=STATUTE_OF_LIMITATIONS_DAYS - 1)
    capped = first < cap
    if capped:
        first = cap
    return first, as_of, (as_of - first).days + 1, capped


def _limitation_note(doc, rows):
    """Пометка юристу: пеня посчитана не с первого дня просрочки, а за последние
    1095 дней (ст. 197 ГК РБ). В претензии это видно в самом расчёте, в иске —
    только по этой пометке."""
    capped = [r for r in rows if r["capped"]]
    if not capped:
        return
    docs = ", ".join(f'{doc_short(r["line"].doc_type)} №{r["line"].doc_number} '
                     f'(срок оплаты {r["line"].due_date.strftime("%d.%m.%Y")})' for r in capped)
    _body(doc, [(f"ВНИМАНИЕ (пометка для юриста, удалить перед подачей): по {docs} с момента "
                 f"просрочки прошло больше трёх лет. Пеня посчитана только за последние "
                 f"{STATUTE_OF_LIMITATIONS_DAYS} дней; срок исковой давности по основному "
                 f"долгу (ст. 197 ГК РБ) проверьте отдельно.", {"bold": True})])


def _period_span(rows, short=True) -> str:
    """«с 28.01.26г. по 18.02.26г. включительно» — от первого дня просрочки
    до последнего расчётного дня. Так подытоживают периоды в ваших исках."""
    first = min(p["с"] for r in rows for p in r["parts"])
    last = max(p["по"] for r in rows for p in r["parts"])
    f = "%d.%m.%y" if short else "%d.%m.%Y"
    return f'с {first.strftime(f)}г. по {last.strftime(f)}г. включительно'


def _all_payments(rows) -> list:
    """Все оплаты по иску без повторов, по возрастанию даты."""
    seen = []
    for r in rows:
        for p in r.get("pays", []):
            if p not in seen:
                seen.append(p)
    return sorted(seen, key=lambda p: p["дата"])


def _shipped(l: DebtLine, pays: list) -> float:
    """Сколько отгружено по накладной: остаток долга плюс уже погашенное.

    В отчёте по ДЗ виден только остаток, а в иске указывается сумма поставки —
    так в вашем образце по машзаводу «Пример».
    """
    return round(l.overdue_amount + sum(p["сумма"] for p in pays), 2)


def _days_word(n: int) -> str:
    """«3 к.дня», «11 к.дней», «22 к.дня» — как в ваших исках."""
    if 11 <= n % 100 <= 14:
        return "дней"
    last = n % 10
    if last == 1:
        return "день"
    if 2 <= last <= 4:
        return "дня"
    return "дней"


def _split(l: DebtLine, as_of: dt.date, pays: list):
    """Части долга по накладной: погашенные и остаток, каждая со своим сроком.

    Погашенная часть считается до дня, предшествующего оплате: день оплаты в
    просрочку не входит. Остаток — до даты документа. Так считает ваш юрист
    (иск по машзаводу «Пример» №01/00001-001).
    """
    first, last, days, capped = _period(l, as_of)
    parts = []
    for p in pays:
        if not (first <= p["дата"] <= as_of):
            continue
        paid_days = (p["дата"] - first).days
        if paid_days <= 0:
            continue
        parts.append(dict(сумма=p["сумма"], с=first, по=p["дата"] - dt.timedelta(days=1),
                          дней=paid_days, оплата=p))
    parts.append(dict(сумма=l.overdue_amount, с=first, по=last, дней=days, оплата=None))
    return parts, capped


def _penalty(lines, rate, as_of, pay_index=None, cp="", contract="", contract_date=""):
    """Пеня по каждой накладной с учётом частичных оплат."""
    rows, debt, penalty = [], 0.0, 0.0
    for l in lines:
        pays = pay.for_line(pay_index or {}, cp, contract, contract_date, l.doc_number, as_of)
        parts, capped = _split(l, as_of, pays)
        for part in parts:
            part["пеня"] = round(part["сумма"] * rate / 100 * part["дней"], 2)
        line_penalty = round(sum(p["пеня"] for p in parts), 2)
        first, last, days, _ = _period(l, as_of)
        debt += l.overdue_amount
        penalty += line_penalty
        rows.append(dict(line=l, first=first, last=last, days=days, capped=capped,
                         amount=line_penalty, parts=parts, pays=pays))
    return rows, round(debt, 2), round(penalty, 2)


def _interest(lines, as_of, pay_index=None, cp="", contract="", contract_date=""):
    rates = load_refi_rates()
    rows, total, no_rate = [], 0.0, 0
    for l in lines:
        pays = pay.for_line(pay_index or {}, cp, contract, contract_date, l.doc_number, as_of)
        parts, _ = _split(l, as_of, pays)
        for part in parts:
            amount, periods, missing = interest_366(part["сумма"], part["с"], part["по"], rates)
            total += amount
            no_rate += missing
            if periods:
                rows.append(dict(line=l, periods=periods, amount=round(amount, 2),
                                 сумма=part["сумма"], оплата=part["оплата"]))
    return rows, round(total, 2), no_rate


def _payments_block(doc, rows):
    """«06.02.26 п/п №101 Ответчиком была произведена оплата в размере …» —
    как в вашем иске по машзаводу «Пример»."""
    seen = []
    for r in rows:
        for p in r.get("pays", []):
            if p in seen:
                continue
            seen.append(p)
            pp = f' п/п №{p["пп"]}' if p["пп"] else ""
            purpose = f' с назначением: {p["назначение"]}' if p["назначение"] else ""
            _body(doc, f'{p["дата"].strftime("%d.%m.%y")}{pp} Ответчиком была произведена '
                       f'оплата в размере {fmt_money(p["сумма"])} руб.{purpose}.')
    return bool(seen)


def _interest_block(doc, rows, total, no_rate):
    _body(doc, "В соответствии со статьей 366 Гражданского кодекса Республики Беларусь за "
               "пользование чужими денежными средствами вследствие их неправомерного удержания, "
               "уклонения от их возврата, иной просрочки в их уплате подлежат уплате проценты на "
               "сумму этих средств. Размер процентов определяется ставкой рефинансирования "
               "Национального банка Республики Беларусь, действовавшей в соответствующие периоды "
               "просрочки.")
    for r in rows:
        l = r["line"]
        base = r.get("сумма", l.overdue_amount)   # часть долга, а не весь остаток
        for p in r["periods"]:
            _body(doc, f'{fmt_money(base)} руб.*{fmt_rate(p["ставка"])}% годовых'
                       f'/360*{p["дней"]} дн. (с {p["с"].strftime("%d.%m.%y")}г. по '
                       f'{p["по"].strftime("%d.%m.%y")}г.) = {fmt_money(p["сумма"])} руб.')
    p = _par(doc, first_line=BODY_INDENT)
    _run(p, f"Итого сумма процентов: {fmt_money(total)} руб.", bold=True)
    if no_rate:
        _body(doc, [(f"ВНИМАНИЕ (пометка для юриста, удалить перед подачей): за {no_rate} дн. "
                     f"ставка рефинансирования неизвестна — дозаполните "
                     f"data/refinancing_rate.csv.", {"bold": True})])


# ---------------------------------------------------------------------------
# ПРЕТЕНЗИЯ — одна форма для прямых и биржевых договоров
# ---------------------------------------------------------------------------

def generate_pretenzia(buyer: BuyerInfo, contract_number: str, contract_date: str,
                       contract_type: str, lines: list, out_path: str,
                       as_of: dt.date = None, deadline_days: int = 10,
                       penalty_rate: float = None):
    as_of = as_of or dt.date.today()
    rate = penalty_rate if penalty_rate is not None else DEFAULT_PENALTY_RATE[contract_type]
    deadline = as_of + dt.timedelta(days=deadline_days)
    dispute = DISPUTE_BODY[contract_type]
    clauses = CONTRACT_CLAUSES[contract_type]

    pay_index = pay.build_index()
    key = (buyer.name, contract_number, contract_date)
    rows, debt, penalty = _penalty(lines, rate, as_of, pay_index, *key)
    if CHARGE_INTEREST_366_IN_PRETENZIA:
        irows, interest, no_rate = _interest(lines, as_of, pay_index, *key)
    else:
        irows, interest, no_rate = [], 0.0, 0
    total = round(debt + penalty + interest, 2)

    doc = Document()
    _setup(doc, left=2.0, right=1.5, top=1.5, bottom=1.5)
    _letterhead(doc)
    _empty(doc)

    _line(doc, f"Исх. № _____ от {ru_date(as_of)}", align=AL.RIGHT)
    _empty(doc)
    _line(doc, buyer.name_full or buyer.name, align=AL.RIGHT, bold=True)
    if buyer.address:
        _line(doc, buyer.address, align=AL.RIGHT)
    if buyer.unp:
        _line(doc, f"УНП {buyer.unp}", align=AL.RIGHT)
    _empty(doc)

    _line(doc, "ПРЕТЕНЗИЯ", align=AL.CENTER, bold=True, size=HEAD_PT - 2)
    _line(doc, "об оплате задолженности и уплате пени", align=AL.CENTER)
    _empty(doc)

    word = "биржевой" if contract_type == "exchange" else "прямой"
    _body(doc, f'Между {SELLER["name_instr"]} (далее — Поставщик) и {buyer.name_full or buyer.name} '
               f'(далее — Покупатель) заключен {word} договор № {contract_number} от '
               f'{contract_date}г. (далее — Договор), в рамках которого Поставщик поставил '
               f'Покупателю товар на условиях отсрочки платежа.')
    _body(doc, "Поставщик надлежащим образом исполнил свои обязательства по Договору, что "
               "подтверждается следующими товарно-транспортными накладными:")
    были_оплаты = False
    for l in lines:
        pays = pay.for_line(pay_index, *key, l.doc_number, as_of)
        были_оплаты = были_оплаты or bool(pays)
        shipped = _shipped(l, pays)
        _body(doc, f'— {doc_full(l.doc_type)} № {l.doc_number} от {l.doc_date} на сумму '
                   f'{fmt_money(shipped)} руб., срок оплаты — '
                   f'{l.due_date.strftime("%d.%m.%Y")}г.')
    указан = "указанной накладной" if len(lines) == 1 else "указанным накладным"
    произведена = ("произведена не в полном объеме" if были_оплаты
                   else "не произведена")
    _body(doc, f'Однако до настоящего времени оплата по {указан} Покупателем '
               f'{произведена}, что является нарушением условий Договора и ст. 290, 456, 486 '
               f'Гражданского кодекса Республики Беларусь, согласно которым обязательства должны '
               f'исполняться надлежащим образом и в установленный срок.')

    p = _par(doc, first_line=BODY_INDENT)
    _run(p, "Расчёт пени за нарушение сроков оплаты:", bold=True)
    for r in rows:
        l = r["line"]
        note = (f' (пеня ограничена {STATUTE_OF_LIMITATIONS_DAYS} дн. в связи со сроком '
                f'исковой давности)') if r["capped"] else ""
        if len(r["parts"]) == 1:
            _body(doc, f'{doc_full(l.doc_type).capitalize()} № {l.doc_number} от {l.doc_date} — '
                       f'{fmt_money(l.overdue_amount)} руб., просрочка {r["days"]} дн. '
                       f'(с {r["first"].strftime("%d.%m.%y")} по {r["last"].strftime("%d.%m.%y")}): '
                       f'{fmt_money(l.overdue_amount)} руб. × {fmt_rate(rate)}% × {r["days"]} дн. = '
                       f'{fmt_money(r["amount"])} руб.{note}')
        else:
            _body(doc, f'{doc_full(l.doc_type).capitalize()} № {l.doc_number} от {l.doc_date} '
                       f'оплачивалась частями, пеня начислена на каждую часть отдельно:')
            for part in r["parts"]:
                tail = ("— остаток" if part["оплата"] is None
                        else f'— погашено {part["оплата"]["дата"].strftime("%d.%m.%y")}'
                             + (f' п/п № {part["оплата"]["пп"]}' if part["оплата"]["пп"] else ""))
                _body(doc, f'{fmt_money(part["сумма"])} руб. × {fmt_rate(rate)}% × '
                           f'{part["дней"]} дн. (с {part["с"].strftime("%d.%m.%y")} по '
                           f'{part["по"].strftime("%d.%m.%y")}) = {fmt_money(part["пеня"])} руб. {tail}')
            _body(doc, f'Итого по накладной: {fmt_money(r["amount"])} руб.{note}')
    p = _par(doc, first_line=BODY_INDENT)
    _run(p, f"Итого сумма пени: {fmt_money(penalty)} руб.", bold=True)

    if irows:
        _interest_block(doc, irows, interest, no_rate)

    parts = [("Таким образом, общая задолженность Покупателя перед Поставщиком составляет "
              f"{fmt_money(total)} руб., из которых {fmt_money(debt)} руб. — сумма основного "
              f"долга, {fmt_money(penalty)} руб. — сумма пени", {"bold": True})]
    if irows:
        parts.append((f", {fmt_money(interest)} руб. — проценты за пользование чужими денежными "
                      f"средствами", {"bold": True}))
    parts.append((f", начисленной по состоянию на {as_of.strftime('%d.%m.%Y')}г.", {"bold": True}))
    _body(doc, parts)

    _body(doc, "На основании изложенного, руководствуясь ст. 290, 311, 456, 486 Гражданского "
               "кодекса Республики Беларусь и условиями Договора,")
    p = _par(doc, first_line=BODY_INDENT)
    _run(p, "Просим:", bold=True)
    _body(doc, f'В течение {deadline_days} (десяти) календарных дней с момента получения '
               f'претензии, то есть не позднее {deadline.strftime("%d.%m.%Y")}г., погасить '
               f'задолженность в размере {fmt_money(total)} руб.')
    _body(doc, f'В случае непогашения задолженности мы будем вынуждены обратиться с исковым '
               f'заявлением в {dispute.get("name_acc", dispute["name"])} согласно пункту договора '
               f'{clauses["dispute_section"]} п.п.{clauses["dispute_clause"]} о взыскании '
               f'образовавшейся задолженности, суммы пени, начисленной на дату фактической '
               f'оплаты, а также расходов, связанных с рассмотрением дела.')

    _empty(doc)
    p = _par(doc, first_line=BODY_INDENT)
    _run(p, "Приложение:", bold=True)
    for l in lines:
        _line(doc, f'— копия {doc_short(l.doc_type)} № {l.doc_number} от {l.doc_date};',
              first_line=BODY_INDENT)
    _line(doc, f'— копия договора № {contract_number} от {contract_date}г.', first_line=BODY_INDENT)

    _signature(doc, executor=False)
    doc.save(out_path)
    return dict(total_debt=debt, total_penalty=penalty, total_interest=interest,
                total_due=total, deadline=deadline)


# ---------------------------------------------------------------------------
# ИСК В ЭКОНОМИЧЕСКИЙ СУД (прямой договор)
# ---------------------------------------------------------------------------

def _isk_ekonsud(doc, buyer, contract_number, contract_date, lines, rate, as_of,
                 rows, debt, penalty, irows, interest, no_rate, total, dispute, clauses, duty):
    _letterhead(doc)
    _empty(doc)

    _addr(doc, dispute["name"])
    _addr(doc, dispute["address"])
    _empty(doc)
    _plaintiff_block(doc)
    _empty(doc)
    _defendant_block(doc, buyer)
    _empty(doc)

    p = _par(doc, left=ADDR_INDENT)
    _run(p, "Цена иска:  ")
    _run(p, f"{fmt_money(total)} руб.", bold=True)
    _empty(doc)
    _line(doc, "Исковое заявление", align=AL.CENTER)
    _empty(doc)

    _body(doc, f'{contract_date_words(contract_date)} между Истцом и Ответчиком был заключен '
               f'договор поставки №{contract_number} (далее – договор), предметом которого '
               f'являлись обязательства Истца поставить Ответчику товар, а Ответчик, в свою '
               f'очередь, обязался принять и произвести оплату стоимости поставленного товара '
               f'в течение {PAYMENT_TERM["direct"]} по факту поставки.')

    doc_list = ", ".join(f'{doc_short(l.doc_type)} №{l.doc_number} от '
                         f'{contract_date_full(l.doc_date)}г.' for l in lines)
    shipped_total = round(sum(_shipped(r["line"], r["pays"]) for r in rows), 2)
    _body(doc, ["Во исполнение обязательств по договору Истец произвел передачу товара Ответчику "
                "на сумму ",
                (f"{fmt_money(shipped_total)} руб.", {"bold": True}),
                f", что подтверждается {doc_list}"])

    dues = sorted({l.due_date for l in lines})
    if len(dues) == 1:
        _body(doc, f'По условиям указанного договора срок оплаты - '
                   f'до {dues[0].strftime("%d.%m.%Y")}г.')
    else:
        _body(doc, "По условиям указанного договора срок оплаты:")
        for l in lines:
            _body(doc, f'по {doc_short(l.doc_type)} №{l.doc_number} от {l.doc_date} — '
                       f'до {l.due_date.strftime("%d.%m.%Y")}г.')
    if any(r["pays"] for r in rows):
        _payments_block(doc, rows)
        _body(doc, "Таким образом, товар оплачен не в полном объеме. Пеня начислена на каждую "
                   "часть задолженности за свой период просрочки.")
    _body(doc, ["Сумма ", ("неоплаченной", {"bold": True}),
                " задолженности по основному долгу по договору составляет ",
                (f"{fmt_money(debt)} руб.", {"bold": True})])

    _body(doc, f'Период просрочки оплаты на {as_of.strftime("%d.%m.%Y")}г. составляет:')
    один = len(rows) == 1 and len(rows[0]["parts"]) == 1
    for r in rows:
        l = r["line"]
        for part in r["parts"]:
            tail = "" if part["оплата"] is None else " (до поступления оплаты)"
            head = "" if один else (f'по {doc_short(l.doc_type)} №{l.doc_number} '
                                    f'от {l.doc_date} на сумму '
                                    f'{fmt_money(part["сумма"])} руб.: ')
            _body(doc, f'{head}с {part["с"].strftime("%d.%m.%y")}г. '
                       f'по {part["по"].strftime("%d.%m.%y")}г. включительно – '
                       f'{part["дней"]} кален. дн.{tail}')
    _limitation_note(doc, rows)

    _body(doc, "В соответствии со статьей 290 Гражданского кодекса Республики Беларусь "
               "обязательства должны исполняться надлежащим образом в соответствии с условиями "
               "обязательства и требованиями законодательства.")
    _body(doc, [f'По условиям п. {clauses["penalty_clause"]} договора:',
                (" " + clauses["penalty_quote"].format(rate=fmt_rate(rate)), {"italic": True})])
    _body(doc, ["В соответствии со статьей 311 Гражданского кодекса Республики Беларусь: ",
                ("«Неустойкой (штрафом, пеней) признается определенная законодательством или "
                 "договором денежная сумма, которую должник обязан уплатить кредитору, если иное "
                 "не предусмотрено законодательными актами, в случае неисполнения или "
                 "ненадлежащего исполнения обязательства, в частности в случае просрочки "
                 "исполнения. По требованию об уплате неустойки кредитор не обязан доказывать "
                 "причинение ему убытков».", {"italic": True})])

    span = _period_span(rows, short=False)
    _body(doc, [("Размер пени", {"bold": True}),
                f", подлежащий уплате Ответчиком за период {span}, составляет ",
                (f"{fmt_money(penalty)} руб.", {"bold": True})])
    for r in rows:
        for part in r["parts"]:
            _body(doc, f'{fmt_money(part["сумма"])} руб.*{fmt_rate(rate)}%*{part["дней"]} дн. '
                       f'просрочки = {fmt_money(part["пеня"])} руб.')

    if irows:
        _interest_block(doc, irows, interest, no_rate)

    parts = ["Таким образом, сумма основного долга составляет ",
             (f"{fmt_money(debt)} руб.", {"bold": True}), ", ",
             (f"{fmt_money(penalty)} руб. ", {"bold": True}),
             f'– сумма пени за просрочку оплаты за период {span}']
    if irows:
        parts += [", ", (f"{fmt_money(interest)} руб. ", {"bold": True}),
                  f'– проценты за пользование чужими денежными средствами по состоянию на '
                  f'{as_of.strftime("%d.%m.%Y")}г.']
    parts += [", а ", ("всего задолженность", {"bold": True}),
              " Ответчика по договору составляет ",
              (f"{fmt_money(total)} руб.", {"bold": True})]
    _body(doc, parts)

    _body(doc, "В целях внесудебного урегулирования хозяйственного конфликта с Ответчиком "
               "неоднократно велись телефонные переговоры о погашении просроченного долга. "
               "Однако, ссылаясь на финансовые затруднения и непоступления оплат от своих "
               "контрагентов, заверения Ответчика об оплате долга по настоящему договору "
               "неоднократно переносились на неопределенный срок.")
    _body(doc, "В соответствии с условиями договора все споры и разногласия, которые могут "
               "возникнуть из настоящего договора или в связи с ним, подлежат рассмотрению в "
               "Экономическом суде по месту нахождения истца.")
    _body(doc, [f'Согласно п.{clauses["dispute_clause"]} договора: ',
                ("«Стороны договорились, что до обращения в Экономический суд с иском по спорам, "
                 "возникающим из настоящего договора, предъявление претензии (письменного "
                 "предложения о добровольном урегулировании спора) либо применение медиации "
                 "не является обязательным.»", {"italic": True})])

    arts = "290, 291, 311, 314, " + ("366, " if irows else "") + "486"
    _body(doc, f'На основании изложенного, руководствуясь ст. {arts} Гражданского кодекса '
               f'Республики Беларусь, ст. 6, 7, 49, 50, 55, 133, 149, 238, 239 Кодекса '
               f'гражданского судопроизводства Республики Беларусь,')
    _line(doc, "ПРОШУ:", first_line=BODY_INDENT)

    parts = [f'Взыскать с {buyer.name_full or buyer.name} в пользу {SELLER["name_gen"]} ',
             (f"{fmt_money(total)} руб.", {"bold": True}),
             ", из которых сумма основного долга составляет ",
             (f"{fmt_money(debt)} руб.", {"bold": True}), ", ",
             (f"{fmt_money(penalty)} руб. ", {"bold": True}),
             f'– сумма пени за просрочку оплаты за период {span}']
    if irows:
        parts += [", ", (f"{fmt_money(interest)} руб. ", {"bold": True}),
                  f'– проценты за пользование чужими денежными средствами по состоянию на '
                  f'{as_of.strftime("%d.%m.%Y")}г.']
    else:
        parts.append(".")
    _body(doc, parts)

    _body(doc, "Истец просит не  применять статью 314 Гражданского кодекса РБ об уменьшении "
               "неустойки так как:")
    _body(doc, "- размер ответственности за просрочку по оплате поставленного товара установлен "
               "соглашением сторон, т.е. Ответчик согласился с установленным договором размером "
               "пени, тем самым добровольно принял на себя обязательство по ее уплате, в случае "
               "нарушения со своей стороны основного обязательства;")
    _body(doc, f'- Ответчику по договору был предоставлен достаточный срок для исполнения '
               f'обязательств по оплате товара – {PAYMENT_TERM["direct"]};')
    _body(doc, "При рассмотрении материалов дела прошу учесть, что в связи с продолжительной "
               "дебиторской задолженностью Ответчика ООО «Ромашка» испытывает недостаток в "
               "собственных денежных средствах и вынуждено обращаться к дорогостоящим заемным "
               "средствам банков.")
    _body(doc, "Прошу  рассмотреть материалы дела  без присутствия Истца и перейти из "
               "подготовительного в основное судебное заседание.")
    if duty["проверено"]:
        _body(doc, f'Размер государственной пошлины составляет {fmt_money(duty["сумма"])} руб. '
                   f'({duty["описание"]}).')
    _body(doc, "Расходы по госпошлине, связанные с разрешением дела в Экономическом суде, прошу "
               "взыскать с Ответчика.")


# ---------------------------------------------------------------------------
# ИСК В АРБИТРАЖНУЮ КОМИССИЮ БУТБ (биржевой договор)
# ---------------------------------------------------------------------------

def _isk_arbitrazh(doc, buyer, contract_number, contract_date, lines, rate, as_of,
                   rows, debt, penalty, irows, interest, no_rate, total, dispute, clauses, duty):
    _letterhead(doc)
    _empty(doc)

    _addr(doc, "Арбитражная комиссия")
    _addr(doc, "ОАО «Белорусская универсальная товарная биржа»")
    _addr(doc, "220 099 г.Минск, ул. Казинца д.2, к.200")
    _empty(doc)
    _plaintiff_block(doc)
    _empty(doc)
    _defendant_block(doc, buyer)
    _empty(doc)

    p = _par(doc, left=ADDR_INDENT)
    _run(p, "Цена иска:")
    _run(p, f" {fmt_money(total)}  руб.", bold=True)
    _empty(doc)
    _line(doc, "Исковое заявление", align=AL.CENTER)
    _empty(doc)

    _body(doc, f'По итогам биржевых торгов ОАО «Белорусская универсальная товарная биржа» '
               f'{contract_date_full(contract_date)}г. между Истцом и Ответчиком был заключен '
               f'договор №{contract_number} (далее – договор), предметом которого являлись '
               f'обязательства Истца передать в собственность Ответчика товар, а Ответчик, '
               f'в свою очередь, обязался произвести оплату стоимости поставленного товара '
               f'в течение {PAYMENT_TERM["exchange"]} по факту поставки.')

    parts = ["Во исполнение обязательств по договору Истец произвел передачу товара Ответчику "]
    for i, r in enumerate(rows):
        l = r["line"]
        parts.append(f'{contract_date_full(l.doc_date)}г. на сумму ' if i == 0
                     else f' и {contract_date_full(l.doc_date)}г. на сумму ')
        parts.append((f'{fmt_money(_shipped(l, r["pays"]))} руб', {"bold": True}))
        parts.append(f'. по {doc_short_ex(l.doc_type)} № {l.doc_number}')
    parts.append(".")
    _body(doc, parts)

    kinds = {doc_short_ex(r["line"].doc_type) for r in rows}
    kind = kinds.pop() if len(kinds) == 1 else "накладным"
    указ = "вышеуказанной" if len(rows) == 1 else "вышеуказанным"
    в_полном = " в полном объеме" if any(r["pays"] for r in rows) else ""
    _body(doc, f'Однако, товар поставленный по {указ} {kind} не оплачен{в_полном} '
               f'до настоящего время.')

    # перечень поступивших оплат — до ссылки на ст. 290, как в образце №154
    _payments_block(doc, rows)

    _body(doc, "В соответствии со статьей 290 Гражданского кодекса Республики Беларусь "
               "обязательства должны исполняться надлежащим образом в соответствии с условиями "
               "обязательства и требованиями законодательства.")

    sents = []
    for r in rows:
        l = r["line"]
        for part in r["parts"]:
            amount_note = ("" if len(r["parts"]) == 1
                           else f' на сумму {fmt_money(part["сумма"])} руб.')
            sents.append(["Период просрочки составил по ",
                          (f"{doc_short_ex(l.doc_type)} №{l.doc_number}", {"bold": True}),
                          f' от {l.doc_date}{amount_note} – {part["дней"]} к. '
                          f'{_days_word(part["дней"])} - с {part["с"].strftime("%d.%m.%y")}г '
                          f'по {part["по"].strftime("%d.%m.%y")}г.'])
    for i, sent in enumerate(sents):
        if i < len(sents) - 1:
            sent[-1] += ","          # в образцах абзацы периодов идут через запятую
        _body(doc, sent)
    _limitation_note(doc, rows)

    _body(doc, f'По условиям п. {clauses["penalty_clause"]} биржевого договора, а также в '
               f'соответствии со статьей 311 Гражданского кодекса Республики Беларусь, размер '
               f'пени, подлежащий уплате Ответчиком, составляет {fmt_rate(rate)}% от '
               f'несвоевременно оплаченной суммы товара за каждый день просрочки оплаты:')
    for r in rows:
        for part in r["parts"]:
            _body(doc, f'{fmt_money(part["сумма"])} руб.*{fmt_rate(rate)}%*{part["дней"]} дн. '
                       f'просрочки = {fmt_money(part["пеня"])} руб.')
    p = _par(doc, first_line=BODY_INDENT)
    _run(p, f"Итого сумма пени: {fmt_money(penalty)} руб.", bold=True)

    if irows:
        _interest_block(doc, irows, interest, no_rate)

    span = _period_span(rows)
    parts = [f'Таким образом, сумма просроченного основного долга по данному договору '
             f'№{contract_number} от {contract_date}г. составляет ',
             (f"{fmt_money(debt)} руб.", {"bold": True}), ", ",
             (f"{fmt_money(penalty)} руб. ", {"bold": True}),
             f'– сумма пени за просрочку оплаты за периоды {span}']
    if irows:
        parts += [", ", (f"{fmt_money(interest)} руб. ", {"bold": True}),
                  f'– проценты за пользование чужими денежными средствами по состоянию на '
                  f'{as_of.strftime("%d.%m.%Y")}г.']
    parts += [", а ", ("всего задолженность", {"bold": True}), " Ответчика составляет ",
              (f"{fmt_money(total)} руб.", {"bold": True})]
    _body(doc, parts)

    _body(doc, f'По условиям п. {clauses["dispute_clause"]} биржевого договора, претензионный '
               f'порядок не является обязательным.')
    _body(doc, "В  соответствии с условиями биржевого договора все споры и разногласия, которые "
               "могут возникнуть из настоящего договора или в связи с ним, подлежат разрешению в "
               "Арбитражной комиссии ОАО «Белорусская универсальная товарная биржа» в "
               "соответствии с регламентом Арбитражной комиссии.")

    arts = "290, 311, " + ("366, " if irows else "") + "456, 476, 486"
    _body(doc, f'На основании изложенного, руководствуясь ст.ст. {arts} Гражданского кодекса '
               f'Республики Беларусь и регламентом Арбитражной комиссии ОАО «Белорусская '
               f'универсальная товарная биржа»,')
    _line(doc, "ПРОШУ:", first_line=BODY_INDENT)

    parts = [f'Взыскать с {buyer.name_full or buyer.name} в пользу {SELLER["name_gen"]} ',
             (f"{fmt_money(total)} руб.", {"bold": True}),
             ", из которых сумма просроченного основного долга составляет ",
             (f"{fmt_money(debt)} руб", {"bold": True}), ", ",
             (f"{fmt_money(penalty)} руб. ", {"bold": True}),
             f'– сумма пени за просрочку оплаты за периоды {span}']
    if irows:
        parts += [", ", (f"{fmt_money(interest)} руб. ", {"bold": True}),
                  f'– проценты за пользование чужими денежными средствами по состоянию на '
                  f'{as_of.strftime("%d.%m.%Y")}г.']
    else:
        parts.append(".")
    _body(doc, parts)

    _body(doc, "При рассмотрении материалов дела прошу учесть, что в связи с длительной "
               "дебиторской задолженностью Ответчика, ООО «Ромашка» испытывает недостаток в "
               "собственных денежных средствах и вынуждено обращаться к дорогостоящим заемным "
               "средствам банков.")
    _body(doc, f'Истец просит не применять статью 314 Гражданского кодекса Республики '
               f'Беларусь. Ответчику был предоставлен достаточный срок для исполнения своих '
               f'обязательств – {PAYMENT_TERM["exchange"]}. Размер пени, предусмотренный '
               f'договором, не высокий ({fmt_rate(rate)}%), а так же это не единственная '
               f'задолженность ответчика.')
    _body(doc, "Прошу  рассмотреть материалы дела  без присутствия Истца.")
    if duty["проверено"]:
        _body(doc, f'Также прошу взыскать с ответчика в полном объеме расходы по уплате '
                   f'арбитражного сбора, связанные с разрешением дела в Арбитражной комиссии, '
                   f'в размере {fmt_money(duty["сумма"])} руб.')
    else:
        _body(doc, "Расходы, связанные с разрешением дела в Арбитражной комиссии, прошу "
                   "взыскать с Ответчика.")


# ---------------------------------------------------------------------------

def generate_isk(buyer: BuyerInfo, contract_number: str, contract_date: str,
                 contract_type: str, lines: list, out_path: str,
                 as_of: dt.date = None, penalty_rate: float = None,
                 court_name: str = "", court_address: str = ""):
    as_of = as_of or dt.date.today()
    rate = penalty_rate if penalty_rate is not None else DEFAULT_PENALTY_RATE[contract_type]
    dispute = dict(DISPUTE_BODY[contract_type])
    if court_name:
        dispute["name"] = court_name
    if court_address:
        dispute["address"] = court_address
    clauses = CONTRACT_CLAUSES[contract_type]
    cfg = ATTACHMENTS[contract_type]

    pay_index = pay.build_index()
    key = (buyer.name, contract_number, contract_date)
    rows, debt, penalty = _penalty(lines, rate, as_of, pay_index, *key)
    if CHARGE_INTEREST_366:
        irows, interest, no_rate = _interest(lines, as_of, pay_index, *key)
    else:
        irows, interest, no_rate = [], 0.0, 0
    total = round(debt + penalty + interest, 2)
    duty = state_duty(total, dispute.get("duty_kind", "исковое_экономсуд"))

    doc = Document()
    if contract_type == "exchange":
        _setup(doc, left=2.22, right=0.68, top=1.5, bottom=0.95)
        _isk_arbitrazh(doc, buyer, contract_number, contract_date, lines, rate, as_of,
                       rows, debt, penalty, irows, interest, no_rate, total, dispute, clauses, duty)
        short = doc_short_ex
    else:
        _setup(doc, left=2.22, right=1.25, top=1.0, bottom=1.25)
        _isk_ekonsud(doc, buyer, contract_number, contract_date, lines, rate, as_of,
                     rows, debt, penalty, irows, interest, no_rate, total, dispute, clauses, duty)
        short = doc_short

    # --- приложения ---------------------------------------------------------
    word = "договора" if contract_type == "exchange" else "договора поставки"
    items = [(f'Копия искового заявления для Ответчика', cfg["isk_sheets"]),
             (f'Копия {word} №{contract_number} от {contract_date_full(contract_date)} г.',
              cfg["contract_sheets"])]
    if cfg["contract_14pt"]:
        items.append((f'Копия {word} №{contract_number} от {contract_date_full(contract_date)}  '
                      f'шрифтом 14 п.т.', cfg["contract_sheets"]))
    for l in lines:
        items.append((f'Копия {short(l.doc_type)} № {l.doc_number} от '
                      f'{contract_date_full(l.doc_date)} г.', cfg["invoice_sheets"]))
    if cfg["pretenzia_copy"]:
        items.append(("Копия претензии и доказательства её направления Ответчику", 1))
    items.extend(cfg["standing"])
    items.append((cfg["duty_doc"], 1))
    for p in _all_payments(rows):
        номер = f' №{p["пп"]}' if p["пп"] else ""
        items.append((f'Копия платежного поручения{номер} от '
                      f'{p["дата"].strftime("%d.%m.%y")}г.', 1))
    _attachments(doc, items)

    _signature(doc, as_of=as_of if contract_type == "exchange" else None,
               executor=(contract_type != "exchange"))

    doc.save(out_path)
    return dict(total_debt=debt, total_penalty=penalty, total_interest=interest,
                total_due=total, duty=duty["сумма"], duty_verified=duty["проверено"])
